# -*- coding: utf-8 -*-
"""把 三航杯赛程.xlsx 转成引擎用的 league JSON。

输入是整理后的表格（见 tools/clean_sanhang_xlsx.py），结构：
    Sheet「甲组赛程」
      A1 标题
      A2 说明
      A4 表头：轮次 | 日期 | 时间 | 主队 | 客队 | 主队进球 | 客队进球 | 状态
      A5+ 45 场比赛

轮次取值：
    数字      已排期的轮次
    「待定」   赛程未公布的比赛，映射为"最后一轮 + 1"

输出：data/sanhang-a.json
"""
import openpyxl, os, json, io
from collections import Counter

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(BASE, 'data', '三航杯赛程.xlsx')
OUT = os.path.join(BASE, 'data', 'sanhang-a.json')

wb = openpyxl.load_workbook(SRC, data_only=True)
sh = wb['甲组赛程']

# ---------- 1. 定位表头 ----------
HEAD_ROW = None
for r in range(1, 20):
    if str(sh.cell(r, 1).value or '').strip() == '轮次':
        HEAD_ROW = r
        break
if HEAD_ROW is None:
    raise SystemExit('找不到表头行（A 列应为「轮次」）')

cols = {}
for c in range(1, sh.max_column + 1):
    name = str(sh.cell(HEAD_ROW, c).value or '').strip()
    if name:
        cols[name] = c

need = ['轮次', '日期', '时间', '主队', '客队', '主队进球', '客队进球', '状态']
missing_cols = [n for n in need if n not in cols]
if missing_cols:
    raise SystemExit('缺少列：' + '、'.join(missing_cols))

# ---------- 2. 读数据行 ----------
rows = []
r = HEAD_ROW + 1
while r <= sh.max_row:
    home = sh.cell(r, cols['主队']).value
    away = sh.cell(r, cols['客队']).value
    if not home or not away:
        r += 1
        continue
    rows.append({
        'round': sh.cell(r, cols['轮次']).value,
        'date': sh.cell(r, cols['日期']).value,
        'time': sh.cell(r, cols['时间']).value,
        'home': str(home).strip(),
        'away': str(away).strip(),
        'hs': sh.cell(r, cols['主队进球']).value,
        'as': sh.cell(r, cols['客队进球']).value,
        'status': str(sh.cell(r, cols['状态']).value or '').strip(),
        'srcRow': r,
    })
    r += 1

if not rows:
    raise SystemExit('没有读到任何比赛行')

# ---------- 3. 轮次映射 ----------
numeric = [x['round'] for x in rows if isinstance(x['round'], (int, float))]
max_round = int(max(numeric)) if numeric else 0
PENDING_ROUND = max_round + 1

def to_round(v):
    if isinstance(v, (int, float)):
        return int(v)
    return PENDING_ROUND      # 「待定」及空值

# ---------- 4. 组装 ----------
teams = sorted({x['home'] for x in rows} | {x['away'] for x in rows})

matches = []
for i, x in enumerate(rows):
    played = x['status'] == '已结束' and x['hs'] is not None and x['as'] is not None
    matches.append({
        'id': 'm%03d' % (i + 1),
        'round': to_round(x['round']),
        'date': str(x['date']).strip() if x['date'] else None,
        'home': x['home'],
        'away': x['away'],
        'homeGoals': int(x['hs']) if played else None,
        'awayGoals': int(x['as']) if played else None,
        'played': played,
    })

pending_count = sum(1 for m in matches if m['round'] == PENDING_ROUND)
played_count = sum(1 for m in matches if m['played'])

league = {
    'id': 'sanhang-a',
    'name': '2026-2027西北工业大学三航杯甲组',
    'season': '2026-2027',
    'teams': [{'id': t, 'name': t} for t in teams],
    'matches': matches,
    'points': {'win': 3, 'draw': 1, 'loss': 0},
    'tiebreak': ['points', 'gd', 'gf', 'wins'],
    'slots': [
        {'from': 1, 'to': 1, 'type': 'title', 'label': '常规赛冠军', 'color': '#ffd54f'},
        {'from': 1, 'to': 4, 'type': 'playoff', 'label': '晋级季后赛', 'color': '#2e7d32'},
        {'from': 9, 'to': 10, 'type': 'relegation', 'label': '降级至乙组', 'color': '#c62828'},
    ],
    'model': {},
    'meta': {
        'source': '三航杯赛程.xlsx',
        # 手工录入：没有"抓取时间"，页面应显示最后一次提交修改时间。
        # 真正的 updatedAt 由 bin/build.js 构建时用 git log 写入。
        'manual': True,
        'maxRound': max_round,
        'note': '校级联赛；10 队单循环 45 场，已排期 %d 场，待排期 %d 场'
                % (len(rows) - pending_count, pending_count),
        'scheduledMatches': len(rows) - pending_count,
        'pendingMatches': pending_count,
    },
}

io.open(OUT, 'w', encoding='utf-8', newline='').write(
    json.dumps(league, ensure_ascii=False, indent=2) + '\n')

# ---------- 5. 校验与报告 ----------
per_team = Counter()
for m in matches:
    per_team[m['home']] += 1
    per_team[m['away']] += 1

rep = []
rep.append('已写入 ' + OUT)
rep.append('  球队：%d 支' % len(teams))
rep.append('  比赛：%d 场（已赛 %d，待排期 %d）' % (len(matches), played_count, pending_count))
rep.append('  已排期轮次：%d，待定轮次：%d' % (max_round, PENDING_ROUND))
rep.append('  每队场次：%s' % ('全部 %d 场' % list(per_team.values())[0]
                              if len(set(per_team.values())) == 1 else dict(per_team)))
rep.append('')
rep.append('各轮场次：')
rc = Counter(m['round'] for m in matches)
for rnd in sorted(rc):
    tag = '（待定）' if rnd == PENDING_ROUND else ''
    rep.append('  轮%-2d %d 场%s' % (rnd, rc[rnd], tag))

io.open(os.path.join(BASE, 'tests', '_json_report.txt'), 'w', encoding='utf-8').write('\n'.join(rep))
print('done')
